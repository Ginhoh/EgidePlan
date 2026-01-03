from openpyxl import load_workbook
from funcoes import *
import customtkinter as ctk
#Dashboard
def exibir_dashboard(position): 
    tabela = load_workbook('total_de_gastos.xlsx')
    verify_sheet(tabela, actual_month())
    main_page = tabela[actual_month()]
    
    #Gastos totais, média, categoria com maior gasto, categoria com menor gasto, comparação com mês anterior
    def gastos_total(planilha):
        total = 0

        for linha in range(2, planilha.max_row+1):
            total += float(planilha[f'B{linha}'].value)
        return total


    def gastos_media(planilha):
        if planilha.max_row -1 == 0:
            return 0
        else:
             return (gastos_total(planilha)) / (planilha.max_row-1)


    def categoria_maior(planilha):
        "Essencial", "Lazer", "Investimentos", "Transporte", "Auto Cuidado"
        essencial = lazer = invest = transp = autoc = 0
        for linha in range(2, planilha.max_row+1):
            if planilha[f'C{linha}'].value == 'Essencial':
                essencial += float(planilha[f'B{linha}'].value)

            elif planilha[f'C{linha}'].value == 'Lazer':
                lazer += float(planilha[f'B{linha}'].value)

            elif planilha[f'C{linha}'].value == 'Investimentos':
                invest += float(planilha[f'B{linha}'].value)
            
            elif planilha[f'C{linha}'].value == 'Transporte':
                transp += float(planilha[f'B{linha}'].value)

            elif planilha[f'C{linha}'].value == 'Auto Cuidado':
                autoc += float(planilha[f'B{linha}'].value)

        valores = {"Essencial": essencial, "Lazer": lazer, "Investimentos": invest, "Transporte": transp, "Auto Cuidado": autoc}
        return max(valores, key=valores.get)
        


    def categoria_menor(planilha):
        essencial = lazer = invest = transp = autoc = 0
        for linha in range(2, planilha.max_row+1):
            if planilha[f'C{linha}'].value == 'Essencial':
                essencial += float(planilha[f'B{linha}'].value)

            elif planilha[f'C{linha}'].value == 'Lazer':
                lazer += float(planilha[f'B{linha}'].value)

            elif planilha[f'C{linha}'].value == 'Investimentos':
                invest += float(planilha[f'B{linha}'].value)
            
            elif planilha[f'C{linha}'].value == 'Transporte':
                transp += float(planilha[f'B{linha}'].value)

            elif planilha[f'C{linha}'].value == 'Auto Cuidado':
                autoc += float(planilha[f'B{linha}'].value)

        valores = {"Essencial": essencial, "Lazer": lazer, "Investimentos": invest, "Transporte": transp, "Auto Cuidado": autoc}
        return min(valores, key=valores.get)

    def comparacao():
        if first(tabela) == False:
            meses = tabela.sheetnames
            tot_anterior = gastos_total(tabela[tabela.sheetnames[len(meses)-2]])
            if tot_anterior > gastos_total(main_page):
                return f'R${tot_anterior - gastos_total(main_page)} a menos que o mês anterior'
            elif tot_anterior < gastos_total(main_page):
                return f'R${gastos_total(main_page) - tot_anterior } a mais que o mês anterior'

    
    new_frame = ctk.CTkFrame(position, width=600, height=600, fg_color='#F3F4F6', border_width=2, border_color='#E5E7EB').place(x=200)
    title = ctk.CTkLabel(new_frame,width=600, text='DASHBOARD',font=('Arial', 20, 'bold'),text_color='#1E3A8A', justify='center', fg_color='#F3F4F6').place(x=200,y=5)


    bloco1=ctk.CTkFrame(position, width=200, height=100, fg_color="#F3F4F6", border_width=2, border_color="#E5E7EB", bg_color='transparent').place(x=230,y=60)
    label_total = ctk.CTkLabel(bloco1, text=f'TOTAL\nR${gastos_total(main_page):.2f}', font=('Arial',24), text_color='#1E3A8A',width=194, height=94, fg_color='#F3F4F6').place(x=232,y=62)

    bloco2=ctk.CTkFrame(position, width=200, height=100, fg_color="#F3F4F6", border_width=2, border_color="#E5E7EB", bg_color='transparent').place(x=570,y=60)
    label_media = ctk.CTkLabel(bloco2, text=f'MÉDIA\nR${gastos_media(main_page):.2f}', font=('Arial',24), text_color='#1E3A8A',width=194, height=94, fg_color='#F3F4F6').place(x=572,y=62)


    label_maior = ctk.CTkLabel(position, text=f'Categoria com maior gasto: {categoria_maior(main_page)}', font=('Arial',20),width=400,anchor='w', text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=230)
    label_menor = ctk.CTkLabel(position, text=f'Categoria com menor  gasto: {categoria_menor(main_page)}', font=('Arial',20),width=400,anchor='w', text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=260)
    label_comparacao = ctk.CTkLabel(position, text=comparacao(), font=('Arial',20), text_color='#1E3A8A', width=400,anchor='w', fg_color='#F3F4F6').place(x=240,y=290)
    