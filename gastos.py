from openpyxl import load_workbook
from datetime import date
from funcoes import *
import customtkinter as ctk


#carregar tabela
table = load_workbook('total_de_gastos.xlsx')
verify_sheet(table, actual_month())
main_page = table[actual_month()]

#mostrar tabela
def show_table(position):

    #Função que abre uma nova janela para exibir tabela anterior
    def last_tb(choice):
            new_window = ctk.CTkToplevel(position, fg_color='#3B82F6')
            new_window.geometry("500x500")
            new_window.title(f"Tabela de {choice}")
            new_window.resizable(False, False)
            new_window.iconbitmap("assents/logo.ico") #Coloca o ícone da aplicação
            new_window.transient(position) # vincula a janela filha à janela mãe
            new_window.grab_set()# impede interação com a janela mãe
            new_window.lift()# traz a janela para frente
            new_scroll = ctk.CTkScrollableFrame(new_window, width=400, height=400, fg_color='#F3F4F6')
            new_scroll.pack(pady=50)
            main_page = table[choice]
            cont = 1
            for linha in range(2, main_page.max_row+1):
                valorA = main_page[f'A{linha}'].value
                valorB = float(main_page[f'B{linha}'].value)
                valorC = main_page[f'C{linha}'].value
                valorD = main_page[f'D{linha}'].value
                new_label = ctk.CTkLabel(new_scroll, text=f'{cont}. {valorA}: R${valorB:.2f} ({valorC}) {valorD}', font=ctk.CTkFont(size=14, weight="bold"), fg_color='#F3F4F6', text_color='#1E3A8A', width=350, anchor='w')
                new_label.pack(fill='x', pady=5)
                cont += 1


    #Exibção da tabela principal
    try:
        cont = 1 
        px = 50
        new_frame = ctk.CTkFrame(position, width=600, height=600, fg_color='#F3F4F6', border_width=2, border_color='#E5E7EB').place(x=200)
        scroll_table = ctk.CTkScrollableFrame(new_frame, width=550, height=400, fg_color='#F3F4F6')
        scroll_table.place(x=230,y=50)
        
        labelGastos = ctk.CTkLabel(position, justify='center', width=600, text="Confira seus gastos aqui", font=ctk.CTkFont(size=20, weight="bold"), fg_color='#F3F4F6', text_color='#1E3A8A').place(x=200,y=10)

       

        for linha in range(2, main_page.max_row+1):
            valorA = main_page[f'A{linha}'].value
            valorB = float(main_page[f'B{linha}'].value)
            valorC = main_page[f'C{linha}'].value
            valorD = main_page[f'D{linha}'].value
            new_label = ctk.CTkLabel(scroll_table, text=f'{cont}. {valorA}: R${valorB:.2f} ({valorC}) {valorD}', font=ctk.CTkFont(size=14, weight="bold"), fg_color='#F3F4F6', text_color='#1E3A8A', width=460, anchor='w',)
            new_label.pack(fill='x', pady=5)

            cont += 1
            px += 30
        btn_add = ctk.CTkButton(position,command=lambda:addGasto(position), width=150, text="Adicionar Gasto", fg_color=('#3B82F6'),text_color='black', border_color='black', border_width=2).place(x=235,y=500)

        btn_remove = ctk.CTkButton(position,command=lambda:remove_gasto(position), width=150, text="Remover Gasto", fg_color=('#3B82F6'),text_color='black', border_color='black', border_width=2).place(x=425,y=500)
        
        btn_fixo = ctk.CTkButton(position,command=lambda:gasto_fixo(position), width=150, text="Criar Gasto Fixo", fg_color=('#3B82F6'),text_color='black', border_color='black', border_width=2).place(x=615,y=500)

        msg_last = ctk.CTkLabel(position, text='Verificar gastos anteriores:',font=ctk.CTkFont(size=14, weight="bold"), fg_color='#F3F4F6', text_color='#1E3A8A').place(x=325,y=535)

        ctg_select = ctk.CTkOptionMenu(position,command=lambda choice: last_tb(choice), width=100,fg_color="#F3F4F6", dropdown_fg_color='#F3F4F6', dropdown_text_color='black', text_color='#3d3d3d', button_color='#F3F4F6', button_hover_color='#575757',
            values= table.sheetnames)
        ctg_select.set("Selecione o mês")
        ctg_select.place(x=525,y=535)
    except:
        new_label = ctk.CTkLabel(new_frame, text=f'Ops! Não foi possível carregar os dados em nosso sistema\n Por favor, tente novamente!', font=ctk.CTkFont(size=12), fg_color='#F3F4F6', text_color='black').place(x=240,y=px)
    finally:   
        pass     
    

#Adicionar gasto
def addGasto(motherWindow):
    #função para enviar os dados para a tabela
    def submitGasto():
        try:
            lastCell = main_page.max_row + 1
            main_page[f'A{lastCell}'].value = descricaoEntry.get()
            valor = valorEntry.get().replace(',','.')
            main_page[f'B{lastCell}'].value = valor

            main_page[f'C{lastCell}'].value =  ctg_select.get()

            today = date.today().strftime('%d/%m/%Y')

            main_page[f'D{lastCell}'].value = today
            labelInfo = ctk.CTkLabel(AddWindow, text="Gasto adicionado com sucesso!", width=400, justify='center', font=ctk.CTkFont(size=16, weight="bold")).place(y=350)

        except ValueError or IndexError or KeyError:
            print('Houve um erro no envio das informações. Revise o que foi pedido e tente novamente.')
        finally:
            show_table(motherWindow)
            table.save('total_de_gastos.xlsx')
            AddWindow.destroy()
            
        
    try:
        AddWindow = ctk.CTkToplevel(motherWindow, fg_color='#F3F4F6')
        AddWindow.geometry("400x400")
        AddWindow.title("Adicionar Gasto")
        AddWindow.resizable(False, False)
        AddWindow.iconbitmap("assents/logo.ico") #Coloca o ícone da aplicação
        AddWindow.transient(motherWindow) # vincula a janela filha à janela mãe
        AddWindow.grab_set()# impede interação com a janela mãe
        AddWindow.lift()# traz a janela para frente

        labelInfo = ctk.CTkLabel(AddWindow, text="Preencha as informações do gasto abaixo:",width=400, justify='center', font=ctk.CTkFont(size=16, weight="bold"), text_color='#1E3A8A').place(x=0,y=20)

        descricaoEntry = ctk.CTkEntry(AddWindow, width=300, placeholder_text="Título do Gasto",fg_color="#E5E7EB", text_color='#3d3d3d')
        descricaoEntry.place(x=50,y=80)
        

        valorEntry = ctk.CTkEntry(AddWindow, width=300, placeholder_text="Valor (Utilize . para as casas decimais.)",fg_color="#E5E7EB", text_color='#3d3d3d')
        valorEntry.place(x=50,y=140)
        


        ctg_select = ctk.CTkOptionMenu(AddWindow, width=300, fg_color="#E5E7EB", dropdown_fg_color='#E5E7EB', dropdown_text_color='black', text_color='#3d3d3d', button_color='#E5E7EB', button_hover_color='#575757',
            values=["Essencial","Alimentação", "Lazer", "Investimentos", "Transporte", "Auto Cuidado"])
        ctg_select.set("Selecione a categoria")
        ctg_select.place(x=50,y=200)
        
        

 
    except ValueError or IndexError or KeyError:
        print('Houve um erro no envio das informações. Revise o que foi pedido e tente novamente.')


    finally:
        btnSubmit = ctk.CTkButton(AddWindow, command=submitGasto, width=150, text="Enviar",text_color='#F3F4F6', fg_color=('#1E3A8A'), hover_color='#3B82F6', border_color='black', border_width=2).place(x=125,y=300)


#Remover gasto
def remove_gasto(motherWindow):
    dialogWindow = ctk.CTkInputDialog(motherWindow,text='Digite qual gasto deseja remover: ', button_fg_color='#1E3A8A', button_hover_color='#3B82F6')
    dialogWindow.geometry("400x200")
    dialogWindow.title("Remover Gasto")
    dialogWindow.iconbitmap("assents/logo.ico") #Coloca o ícone da aplicação
    value_to_remove = dialogWindow.get_input()
    if value_to_remove is None:
        return
    value_to_remove = value_to_remove.strip()
    if value_to_remove == '':
        return
    if value_to_remove != '':
        main_page.delete_rows(int(value_to_remove)+1)
        table.save('total_de_gastos.xlsx')
        show_table(motherWindow)


def gasto_fixo(motherWindow):
    try:
        AddWindow = ctk.CTkToplevel(motherWindow, fg_color='#F3F4F6')
        AddWindow.geometry("400x400")
        AddWindow.title("Criar Gasto Fixo")
        AddWindow.resizable(False, False)
        AddWindow.iconbitmap("assents/logo.ico") #Coloca o ícone da aplicação
        AddWindow.transient(motherWindow) # vincula a janela filha à janela mãe
        AddWindow.grab_set()# impede interação com a janela mãe
        AddWindow.lift()# traz a janela para frente

        labelInfo = ctk.CTkLabel(AddWindow, text="Preencha as informações do gasto abaixo:",width=400, justify='center', font=ctk.CTkFont(size=16, weight="bold"), text_color='#1E3A8A').place(x=0,y=20)

        descricaoEntry = ctk.CTkEntry(AddWindow, width=300, placeholder_text="Título do Gasto",fg_color="#E5E7EB", text_color='#3d3d3d')
        descricaoEntry.place(x=50,y=80)
        

        valorEntry = ctk.CTkEntry(AddWindow, width=300, placeholder_text="Valor (Utilize . para as casas decimais.)",fg_color="#E5E7EB", text_color='#3d3d3d')
        valorEntry.place(x=50,y=140)
        


        ctg_select = ctk.CTkOptionMenu(AddWindow, width=300, fg_color="#E5E7EB", dropdown_fg_color='#E5E7EB', dropdown_text_color='black', text_color='#3d3d3d', button_color='#E5E7EB', button_hover_color='#575757',
            values=["Essencial","Alimentação", "Lazer", "Investimentos", "Transporte", "Auto Cuidado"])
        ctg_select.set("Selecione a categoria")
        ctg_select.place(x=50,y=200)

    except ValueError or IndexError or KeyError:
        print('Houve um erro no envio das informações. Revise o que foi pedido e tente novamente.')

    finally:
        btnSubmit = ctk.CTkButton(AddWindow, width=150, text="Criar Gasto",text_color='#F3F4F6', fg_color=('#1E3A8A'), hover_color='#3B82F6', border_color='black', border_width=2).place(x=125,y=300)
    