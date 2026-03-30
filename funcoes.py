from openpyxl import workbook, load_workbook
import customtkinter as ctk

#retornar mês atual
def actual_month():
    from datetime import date
    mes = date.today().month
    if mes == 1: return '01'
    elif mes == 2: return '02'
    elif mes == 3: return '03'
    elif mes == 4: return '04'
    elif mes == 5: return '05'
    elif mes == 6: return '06'
    elif mes == 7: return '07'
    elif mes == 8: return '08'
    elif mes == 9: return '09'
    elif mes == 10: return '10'
    elif mes == 11: return '11'
    elif mes == 12: return '12'


#criar_abas():
def verify_sheet(table, nameAba):
    if nameAba not in table.sheetnames:
        table.create_sheet(nameAba)

        table[nameAba]['A1'].value = 'Descrição'
        table[nameAba]['B1'].value = 'Valor'
        table[nameAba]['C1'].value = 'Tipo de gasto'
        table[nameAba]['D1'].value = 'Data'
    if 'Sheet' in table.sheetnames:
        table.remove(table['Sheet'])
    if 'Plan1' in table.sheetnames:
        table.remove(table['Plan1'])

#verificar se é a primeira aba
def first(tabela):
    cont = 0
    for aba in tabela.sheetnames:
        cont += 1
    if cont == 1:
        return True
    else:
        return False  
    


