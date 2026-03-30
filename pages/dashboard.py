from openpyxl import load_workbook
from funcoes import *
import customtkinter as ctk
#Dashboard



def exibir_dashboard(position): 
    tabela = load_workbook('total_de_gastos.xlsx')
    main_page = tabela[actual_month()]

    new_frame = ctk.CTkFrame(position, width=600, height=600, fg_color='#F3F4F6', border_width=2, border_color='#E5E7EB').place(x=200)
    title = ctk.CTkLabel(new_frame,width=600, text='DASHBOARD',font=('Arial', 20, 'bold'),text_color='#1E3A8A', justify='center', fg_color='#F3F4F6').place(x=200,y=5)


    if first(tabela) == True:
        ctg_select = ctk.CTkOptionMenu(new_frame, width=100,command=lambda choice: last_table(choice),fg_color="#F3F4F6", dropdown_fg_color='#F3F4F6', dropdown_text_color='black', text_color='#3d3d3d', button_color='#F3F4F6', button_hover_color='#575757',
            values= tabela.sheetnames)
        ctg_select.set("Selecione o mês")
        ctg_select.place(x=630,y=5)
        


    bloco1=ctk.CTkFrame(position, width=200, height=100, fg_color="#F3F4F6", border_width=2, border_color="#E5E7EB", bg_color='transparent').place(x=230,y=60)
    label_total = ctk.CTkLabel(bloco1, text=f'TOTAL\nR${gastos_total(main_page):.2f}', font=('Arial',24), text_color='#1E3A8A',width=194, height=94, fg_color='#F3F4F6').place(x=232,y=62)

    bloco2=ctk.CTkFrame(position, width=200, height=100, fg_color="#F3F4F6", border_width=2, border_color="#E5E7EB", bg_color='transparent').place(x=570,y=60)
    label_media = ctk.CTkLabel(bloco2, text=f'MÉDIA\nR${gastos_media(main_page):.2f}', font=('Arial',24), text_color='#1E3A8A',width=194, height=94, fg_color='#F3F4F6').place(x=572,y=62)


    label_maior = ctk.CTkLabel(position, text=f'Categoria com maior gasto: {categoria_maior(main_page)}', font=('Arial',20),width=400,anchor='w', text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=230)

    label_menor = ctk.CTkLabel(position, text=f'Categoria com menor  gasto: {categoria_menor(main_page)}', font=('Arial',20),width=400,anchor='w', text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=260)

    label_comparacao = ctk.CTkLabel(position, text=comparacao(), font=('Arial',20), text_color='#1E3A8A', width=400,anchor='w', fg_color='#F3F4F6').place(x=240,y=290)

    subtitle = ctk.CTkLabel(position,width=600, text='Gastos por Categoria',font=('Arial', 16, 'bold'),text_color='#1E3A8A', justify='center', fg_color='#F3F4F6').place(x=200,y=320)

    label_essencial = ctk.CTkLabel(position, text=f'Essencial: R${gastos_por_categoria(main_page, "Essencial"):.2f}', font=('Arial',14), text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=360)

    label_alimentacao = ctk.CTkLabel(position, text=f'Alimentação: R${gastos_por_categoria(main_page, "Alimentação"):.2f}', font=('Arial',14), text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=390)

    label_lazer = ctk.CTkLabel(position, text=f'Lazer: R${gastos_por_categoria(main_page, "Lazer"):.2f}', font=('Arial',14), text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=420)
    label_investimentos = ctk.CTkLabel(position, text=f'Investimentos: R${gastos_por_categoria(main_page, "Investimentos"):.2f}', font=('Arial',14), text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=450)

    label_transporte = ctk.CTkLabel(position, text=f'Transporte: R${gastos_por_categoria(main_page, "Transporte"):.2f}', font=('Arial',14), text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=480)
    
    label_autocuidado = ctk.CTkLabel(position, text=f'Auto-Cuidado: R${gastos_por_categoria(main_page, "Auto Cuidado"):.2f}', font=('Arial',14), text_color='#1E3A8A', fg_color='#F3F4F6').place(x=240,y=510)
    


    def last_table(mouth):
        main_page = tabela[mouth]
        exibir_dashboard(position)
        print(main_page)


#Gastos totais, média, categoria com maior gasto, categoria com menor gasto, comparação com mês anterior
def gastos_total(planilha):
    total = 0

    for linha in range(2, planilha.max_row+1):
        total += float(planilha[f'B{linha}'].value)
    return total


def gastos_media(planilha):
    if planilha.max_row -1 == 0:
        return 0
    else:
            return (gastos_total(planilha)) / (planilha.max_row-1)


def categoria_maior(planilha):
    "Essencial", "Lazer", "Investimentos", "Transporte", "Auto Cuidado"
    essencial = lazer = invest = transp = autoc = 0
    for linha in range(2, planilha.max_row+1):
        if planilha[f'C{linha}'].value == 'Essencial':
            essencial += float(planilha[f'B{linha}'].value)

        elif planilha[f'C{linha}'].value == 'Lazer':
            lazer += float(planilha[f'B{linha}'].value)

        elif planilha[f'C{linha}'].value == 'Investimentos':
            invest += float(planilha[f'B{linha}'].value)
        
        elif planilha[f'C{linha}'].value == 'Transporte':
            transp += float(planilha[f'B{linha}'].value)

        elif planilha[f'C{linha}'].value == 'Auto Cuidado':
            autoc += float(planilha[f'B{linha}'].value)

    valores = {"Essencial": essencial, "Lazer": lazer, "Investimentos": invest, "Transporte": transp, "Auto Cuidado": autoc}
    return max(valores, key=valores.get)
    


def categoria_menor(planilha):
    essencial = lazer = invest = transp = autoc = 0
    for linha in range(2, planilha.max_row+1):
        if planilha[f'C{linha}'].value == 'Essencial':
            essencial += float(planilha[f'B{linha}'].value)

        elif planilha[f'C{linha}'].value == 'Lazer':
            lazer += float(planilha[f'B{linha}'].value)

        elif planilha[f'C{linha}'].value == 'Investimentos':
            invest += float(planilha[f'B{linha}'].value)
        
        elif planilha[f'C{linha}'].value == 'Transporte':
            transp += float(planilha[f'B{linha}'].value)

        elif planilha[f'C{linha}'].value == 'Auto Cuidado':
            autoc += float(planilha[f'B{linha}'].value)

    valores = {"Essencial": essencial, "Lazer": lazer, "Investimentos": invest, "Transporte": transp, "Auto Cuidado": autoc}
    return min(valores, key=valores.get)

def comparacao():
    tabela = load_workbook('total_de_gastos.xlsx')
    main_page = tabela[actual_month()]
    if first(tabela) == False:
        meses = tabela.sheetnames
        tot_anterior = gastos_total(tabela[tabela.sheetnames[len(meses)-2]])
        if tot_anterior > gastos_total(main_page):
            return f'R${tot_anterior - gastos_total(main_page):.2f} a menos que o mês anterior'
        elif tot_anterior < gastos_total(main_page):
            return f'R${gastos_total(main_page) - tot_anterior:.2f} a mais que o mês anterior'
        

def gastos_por_categoria(planilha, categoria):
    total_categoria = 0
    for linha in range(2, planilha.max_row+1):
        if planilha[f'C{linha}'].value == categoria:
            total_categoria += float(planilha[f'B{linha}'].value)
    return total_categoria


