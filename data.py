#DashBoard estatístico
from openpyxl import load_workbook
from datetime import date
from funcoes import *
from time import sleep
import customtkinter as ctk



tabela = load_workbook('total_de_gastos.xlsx')
main_page = tabela[mes_atual()]


verificarAba(tabela, mes_atual())


#main_page.delete_rows(6)# - Serve para apagar linhas
def exibir_tabela(position):
    try:
        valores = cont = 0 
        px = 50
        new_frame = ctk.CTkFrame(position, width=600, height=400, fg_color='#F3F4F6',border_color='#E5E7EB', border_width=2).place(x=200)
        for linha in range(2, main_page.max_row+1):
            sleep(0.5)
            valorA = main_page[f'A{linha}'].value
            valorB = main_page[f'B{linha}'].value
            valorC = main_page[f'C{linha}'].value
            valorD = main_page[f'D{linha}'].value
            new_label = ctk.CTkLabel(new_frame, text=f'{valorA}: R${valorB} ({valorC}) {valorD}', font=ctk.CTkFont(size=12), fg_color='#F3F4F6', text_color='#4A6D7C').place(x=240,y=px)
            num = float(valorB)
            valores += num
            cont += 1
            px += 30
    except:
        new_label = ctk.CTkLabel(new_frame, text=f'Ops! Não foi possível carregar os dados em nosso sistema\n Por favor, tente novamente!', font=ctk.CTkFont(size=12), fg_color='#F3F4F6', text_color='black').place(x=240,y=px)
    finally:   
        total_label = ctk.CTkLabel(new_frame, text=f'Foram exibidos {cont} itens.\nValor gasto total: R${valores:.2f}', font=ctk.CTkFont(size=14, weight="bold"), fg_color='#F3F4F6', text_color='black').place(x=240,y=px+40)
        
    


def addGasto(motherWindow):
    def submitGasto():
        try:
            lastCell = main_page.max_row + 1
            main_page[f'A{lastCell}'].value = descricaoEntry.get()

            main_page[f'B{lastCell}'].value = valorEntry.get()

            main_page[f'C{lastCell}'].value =  criterioEntry.get()

            today = date.today().strftime('%d/%m/%Y')

            main_page[f'D{lastCell}'].value = today
            labelInfo = ctk.CTkLabel(AddWindow, text="Gasto adicionado com sucesso!", width=400, justify='center', font=ctk.CTkFont(size=16, weight="bold")).place(y=350)

        except ValueError or IndexError or KeyError:
            print('Houve um erro no envio das informações. Revise o que foi pedido e tente novamente.')
        finally:
            descricaoEntry.delete(0, 'end') #Limpa o campo após o envio
            valorEntry.delete(0, 'end')
            criterioEntry.delete(0, 'end')
            motherWindow.update()
            exibir_tabela(motherWindow)
            AddWindow.destroy()
            tabela.save('total_de_gastos.xlsx')
            
        
    try:
        AddWindow = ctk.CTkToplevel(motherWindow, fg_color='#F3F4F6')
        AddWindow.geometry("400x400")
        AddWindow.title("Adicionar Gasto")
        AddWindow.resizable(False, False)
        AddWindow.iconbitmap("assents/logo.ico") #Coloca o ícone da aplicação
        AddWindow.transient(motherWindow)
        AddWindow.grab_set()# impede interação com a janela mãe
        AddWindow.lift()# traz a janela para frente

        labelInfo = ctk.CTkLabel(AddWindow, text="Preencha as informações do gasto abaixo:",width=400, justify='center', font=ctk.CTkFont(size=16, weight="bold"), text_color='#1E3A8A').place(x=0,y=20)

        descricaoEntry = ctk.CTkEntry(AddWindow, width=300, placeholder_text="Título do Gasto")
        descricaoEntry.place(x=50,y=80)
        

        valorEntry = ctk.CTkEntry(AddWindow, width=300, placeholder_text="Valor (Utilize . para as casas decimais.)")
        valorEntry.place(x=50,y=140)
        

        criterioEntry = ctk.CTkEntry(AddWindow, width=300, placeholder_text="Grau de importância")
        criterioEntry.place(x=50,y=200)
        

 
    except ValueError or IndexError or KeyError:
        print('Houve um erro no envio das informações. Revise o que foi pedido e tente novamente.')


    finally:
        btnSubmit = ctk.CTkButton(AddWindow, command=submitGasto, width=150, text="Enviar", fg_color=('#1E3A8A'),text_color='black', border_color='black', border_width=2).place(x=125,y=300)

        

    #Se salvar com um nome diferente, ele cria um arquivo

#main_page.max_column #Ver o máximo de colunas
#main_page.max_row #Ver o máximo de linhas

#main_page['A1'].value é possível exibir e alterar o valor de uma célulam