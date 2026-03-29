import sqlite3
from openpyxl import load_workbook

table = load_workbook('total_de_gastos.xlsx')
main_page = table[table.sheetnames[-1]]

db = sqlite3.connect('gastos.db')
cursor = db.cursor()

cursor.execute("""CREATE TABLE IF NOT EXISTS gastos (
               id INTEGER PRIMARY KEY AUTOINCREMENT,
               nome TEXT(100) NOT NULL,
               valor REAL NOT NULL,
               categoria TEXT(50) NOT NULL,
               data DATE NOT NULL
               )""")

cursor.execute("""CREATE TABLE IF NOT EXISTS fixos (
               id INTEGER PRIMARY KEY AUTOINCREMENT,
               nome TEXT(100) NOT NULL,
               valor REAL NOT NULL,
               categoria TEXT(50) NOT NULL,
               data DATE NOT NULL
               )""")


# cursor.execute("""SELECT nome, valor, categoria, data FROM gastos""")
# query = cursor.fetchall()
# for dados in query:
#     valorA, valorB, valorC,valorD = dados
#     print(f'{valorA}: R${valorB:.2f} ({valorC}) {valorD}')



# for linha in range(2, main_page.max_row+1):
#     valorA = main_page[f'A{linha}'].value
#     valorB = float(main_page[f'B{linha}'].value)
#     valorC = main_page[f'C{linha}'].value
#     valorD = main_page[f'D{linha}'].value

#     cursor.execute("""INSERT INTO GASTOS (nome, valor, categoria, data)
#                    VALUES (?, ?, ?, ?)""", (valorA, valorB, valorC, valorD))
# for i in table.sheetnames:
#     page = table[i]
#     for linha in range(2, page.max_row+1):
#         valorA = page[f'A{linha}'].value
#         valorB = float(page[f'B{linha}'].value)
#         valorC = page[f'C{linha}'].value
#         valorD = page[f'D{linha}'].value

#         cursor.execute("""INSERT INTO GASTOS (nome, valor, categoria, data)
#                        VALUES (?, ?, ?, ?)""", (valorA, valorB, valorC, valorD))
        

#cursor.execute("""DELETE FROM gastos""") #Limpa a tabela para evitar duplicação de dados

db.commit()
